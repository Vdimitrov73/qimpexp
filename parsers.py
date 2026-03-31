"""
parsers.py - Workbook parsing, config loading, account/security resolution.

Key behaviour:
  - ROC transactions ALWAYS use the running share balance (col G), never col D.
    If col D is populated on a ROC row a warning is emitted but the value is ignored.
  - Col G (Share Balance) is read directly when present and used as the authoritative
    share balance for ROC amount computation (AI feedback item 2).
  - Negative share balances trigger a warning (AI feedback item 7).
  - Zero-share ROC rows are skipped with a warning (AI feedback item 8).
"""

import json
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    import openpyxl
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

SUPPORTED_TYPES = {"Buy", "Sell", "ROC"}
EXCEL_EPOCH = date(1899, 12, 30)

_DATE_FORMATS = [
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%d/%m/%Y",
    "%Y/%m/%d",
    "%Y-%b-%d",   # e.g. 2025-Apr-01 (Excel formula text)
    "%Y-%B-%d",
    "%d-%b-%Y",
    "%d-%B-%Y",
]


# ---------------------------------------------------------------------------
# Date / Decimal helpers
# ---------------------------------------------------------------------------

def _excel_date_to_python(serial):
    try:
        n = int(float(str(serial).strip()))
        if n <= 0:
            return None
        if n >= 60:
            n -= 1
        return EXCEL_EPOCH + timedelta(days=n)
    except (ValueError, OverflowError, TypeError):
        return None


def _parse_date_value(raw):
    """Parse any cell value (datetime, date, string, serial) to a Python date."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return _excel_date_to_python(s)


def _parse_decimal(raw):
    if raw is None or str(raw).strip() == "":
        return None
    try:
        s = (str(raw).strip()
             .replace(",", "")
             .replace("$", "")
             .replace("(", "-")
             .replace(")", ""))
        return Decimal(s)
    except InvalidOperation:
        return None


# ---------------------------------------------------------------------------
# Single-row parser
# ---------------------------------------------------------------------------

def _parse_row_values(raw_date, raw_type, raw_price, raw_shares, raw_commission,
                      raw_col_g,   # column G: Share Balance (may be None or a formula result)
                      ticker, row_num, warnings):
    """
    Parse pre-extracted cell values into a raw transaction dict.

    For ROC rows:
      - raw_shares (col D) is intentionally ignored for amount computation.
      - raw_col_g (col G) is stored as 'col_g_balance' for use in
        reconstruct_share_balances().
      - A warning is emitted when col D is non-empty on a ROC row.
    """
    type_str = str(raw_type).strip() if raw_type is not None else ""

    if type_str.lower() in ("", "type", "transaction", "transaction type", "action"):
        return None, None

    type_upper = type_str.upper()
    if type_upper == "ROC":
        type_norm = "ROC"
    elif type_upper == "BUY":
        type_norm = "Buy"
    elif type_upper == "SELL":
        type_norm = "Sell"
    else:
        warn = f"[WARN] {ticker} row {row_num}: unsupported type {type_str!r} — skipped"
        return None, warn

    trade_date = _parse_date_value(raw_date)
    if trade_date is None:
        date_display = repr(str(raw_date)) if raw_date is not None else "''"
        warn = f"[WARN] {ticker} row {row_num}: unparseable date {date_display} — skipped"
        return None, warn

    price      = _parse_decimal(raw_price)
    shares_col_d = _parse_decimal(raw_shares)    # col D — only used for Buy/Sell
    commission = _parse_decimal(raw_commission) or Decimal("0")
    col_g_bal  = _parse_decimal(raw_col_g)       # col G — authoritative balance for ROC

    if type_norm in ("Buy", "Sell"):
        if price is None:
            warn = f"[WARN] {ticker} row {row_num}: missing price for {type_norm} — skipped"
            return None, warn
        if shares_col_d is None:
            warn = f"[WARN] {ticker} row {row_num}: missing shares for {type_norm} — skipped"
            return None, warn

    if type_norm == "ROC":
        if price is None:
            warn = f"[WARN] {ticker} row {row_num}: missing ROC per-unit value — skipped"
            return None, warn
        # Warn (but don't fail) when col D is populated on a ROC row —
        # it might be a user data-entry mistake.
        if shares_col_d is not None:
            warnings.append(
                f"[WARN] {ticker} row {row_num}: ROC row has shares in col D "
                f"({shares_col_d}) — col D ignored, share balance used instead"
            )

    if type_norm == "Buy":
        normalized_action = "Buy"
        total_amount = (price * shares_col_d + commission).quantize(Decimal("0.01"))
        memo = None
    elif type_norm == "Sell":
        normalized_action = "Sell"
        total_amount = (price * shares_col_d - commission).quantize(Decimal("0.01"))
        memo = None
    else:  # ROC — total_amount computed later in reconstruct_share_balances
        normalized_action = "RtrnCapX"
        total_amount = None
        if price < 0:
            memo = "Phantom distribution"
        else:
            memo = "ROC adjustment"

    return {
        "source_sheet":       ticker,
        "source_row_number":  row_num,
        "ticker":             ticker,
        "trade_date":         trade_date,
        "source_type":        type_str,
        "normalized_action":  normalized_action,
        "price_per_share":    price,
        "shares":             shares_col_d if type_norm in ("Buy", "Sell") else None,
        "commission":         commission,
        "col_g_balance":      col_g_bal,    # col G value stored for ROC use
        "derived_share_balance": None,
        "total_amount":       total_amount,
        "memo":               memo,
        "account_name":       None,
        "security_name":      None,
    }, None


# ---------------------------------------------------------------------------
# Share balance reconstruction
# ---------------------------------------------------------------------------

def reconstruct_share_balances(transactions, verbose=False):
    """
    Compute running share balance per ticker in chronological order.

    ROC amount resolution priority:
      1. col_g_balance from the workbook (authoritative pre-computed value)
      2. Running balance accumulated from Buy/Sell rows above this point

    A warning is emitted (and the row skipped) if the resolved share count
    is zero — that would produce a meaningless zero-value ROC entry.
    A warning is emitted if the running balance goes negative after a Sell.
    """
    from collections import defaultdict
    warnings_out = []

    by_ticker = defaultdict(list)
    for t in transactions:
        by_ticker[t["ticker"]].append(t)

    result = []
    for ticker, txns in by_ticker.items():
        txns_sorted = sorted(txns, key=lambda x: (x["trade_date"], x["source_row_number"] or 0))
        running_balance = Decimal("0")

        for txn in txns_sorted:
            txn = dict(txn)
            action = txn["normalized_action"]

            if action == "Buy":
                running_balance += txn["shares"]

            elif action == "Sell":
                running_balance -= txn["shares"]
                if running_balance < 0:
                    w = (f"[WARN] {ticker} row {txn['source_row_number']}: "
                         f"share balance went negative ({running_balance}) after Sell")
                    warnings_out.append(w)
                    if verbose:
                        print(w)

            elif action == "RtrnCapX":
                # Priority 1: use col G balance from workbook (pre-computed, authoritative)
                col_g = txn.get("col_g_balance")
                if col_g is not None and col_g > 0:
                    shares_for_roc = col_g
                    if verbose:
                        print(f"[BALANCE] {ticker} row {txn['source_row_number']}: "
                              f"ROC using col G balance {col_g}")
                else:
                    if col_g is not None and col_g < 0:
                        w = (f"[WARN] {ticker} row {txn['source_row_number']}: "
                             f"col G balance is negative ({col_g}) — "
                             f"falling back to running balance")
                        warnings_out.append(w)
                        if verbose:
                            print(w)
                    shares_for_roc = running_balance
                    if verbose:
                        print(f"[BALANCE] {ticker} row {txn['source_row_number']}: "
                              f"ROC shares from running balance {running_balance}")

                # Zero-share guard
                if shares_for_roc == 0:
                    w = (f"[WARN] {ticker} row {txn['source_row_number']}: "
                         f"ROC share balance is zero — skipped")
                    warnings_out.append(w)
                    if verbose:
                        print(w)
                    # Don't append this txn to result
                    continue

                txn["shares"] = shares_for_roc
                per_unit = txn["price_per_share"]
                if per_unit is not None:
                    txn["total_amount"] = (per_unit * shares_for_roc).quantize(Decimal("0.01"))
                else:
                    txn["total_amount"] = Decimal("0.00")

            txn["derived_share_balance"] = running_balance
            result.append(txn)

    result.sort(key=lambda x: (x["trade_date"], x["source_row_number"] or 0))
    return result, warnings_out


# ---------------------------------------------------------------------------
# Main workbook loader
# ---------------------------------------------------------------------------

def load_workbook_transactions(acb_path, fund_filter=None, verbose=False):
    """
    Open the .xlsx workbook and parse all matching sheets.
    Returns dict with: transactions, warnings, sheets_found, fatal.
    """
    warnings = []
    fund_filter_upper = {f.strip().upper() for f in fund_filter} if fund_filter else None

    if not Path(acb_path).exists():
        return {"transactions": [], "warnings": [], "sheets_found": [],
                "fatal": f"File not found: {acb_path}"}

    if not _HAS_OPENPYXL:
        return {"transactions": [], "warnings": [], "sheets_found": [],
                "fatal": "openpyxl is required. Install: pip install openpyxl"}

    try:
        wb = openpyxl.load_workbook(str(acb_path), read_only=True, data_only=True)
    except Exception as exc:
        return {"transactions": [], "warnings": [], "sheets_found": [],
                "fatal": f"Cannot open workbook: {exc}"}

    transactions = []
    sheets_found = []

    for sheet_name in wb.sheetnames:
        ticker = sheet_name.strip().upper()

        if fund_filter_upper and ticker not in fund_filter_upper:
            if verbose:
                print(f"[SHEET] Skipping {sheet_name!r} (not in fund filter)")
            continue

        ws = wb[sheet_name]
        sheet_txns = 0
        row_num = 1

        for row in ws.iter_rows(min_row=2, values_only=True):
            row_num += 1
            if all(cell is None for cell in (row[:5] if len(row) >= 5 else row)):
                continue

            raw_date       = row[0] if len(row) > 0 else None
            raw_type       = row[1] if len(row) > 1 else None
            raw_price      = row[2] if len(row) > 2 else None
            raw_shares     = row[3] if len(row) > 3 else None  # col D
            raw_commission = row[4] if len(row) > 4 else None  # col E
            # col F (index 5) = Capital Gains — skip
            raw_col_g      = row[6] if len(row) > 6 else None  # col G = Share Balance

            txn, warn = _parse_row_values(
                raw_date, raw_type, raw_price, raw_shares, raw_commission,
                raw_col_g, ticker, row_num, warnings
            )
            if txn is not None:
                transactions.append(txn)
                sheet_txns += 1
            elif warn:
                warnings.append(warn)

        if verbose:
            print(f"[SHEET] {sheet_name!r} (ticker={ticker}): {sheet_txns} transactions")

        if sheet_txns > 0:
            sheets_found.append(sheet_name)

    wb.close()

    if not sheets_found:
        return {"transactions": [], "warnings": warnings, "sheets_found": [],
                "fatal": "No valid sheets with transactions found in workbook."}

    return {"transactions": transactions, "warnings": warnings,
            "sheets_found": sheets_found, "fatal": None}


# ---------------------------------------------------------------------------
# Date filter
# ---------------------------------------------------------------------------

def apply_date_filter(transactions, start_date, end_date, verbose=False):
    if start_date is None and end_date is None:
        return transactions
    filtered = []
    for t in transactions:
        d = t["trade_date"]
        if start_date and d < start_date:
            if verbose:
                print(f"[FILTER] {t['ticker']} row {t['source_row_number']} "
                      f"date {d} < start {start_date} — excluded")
            continue
        if end_date and d > end_date:
            if verbose:
                print(f"[FILTER] {t['ticker']} row {t['source_row_number']} "
                      f"date {d} > end {end_date} — excluded")
            continue
        filtered.append(t)
    return filtered


# ---------------------------------------------------------------------------
# account_periods.json
# ---------------------------------------------------------------------------

def load_account_periods(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Malformed account_periods.json: {exc}") from exc

    if not isinstance(data, dict):
        raise ValueError("account_periods.json must be a JSON object.")

    parsed = {}
    for ticker, periods in data.items():
        if not isinstance(periods, list):
            raise ValueError(f"account_periods.json: {ticker!r} must be a list.")
        parsed_periods = []
        for p in periods:
            account   = p.get("account")
            start_raw = p.get("start")
            end_raw   = p.get("end")
            if not account or not start_raw:
                raise ValueError("Each period needs 'account' and 'start'.")
            start_dt = datetime.strptime(start_raw, "%Y-%m-%d").date()
            end_dt   = datetime.strptime(end_raw,   "%Y-%m-%d").date() if end_raw else None
            parsed_periods.append({"account": account, "start": start_dt, "end": end_dt})
        parsed_periods.sort(key=lambda x: x["start"])
        for i in range(len(parsed_periods) - 1):
            a, b = parsed_periods[i], parsed_periods[i + 1]
            if a["end"] is None:
                raise ValueError(
                    f"account_periods.json: {ticker!r} has an open-ended period "
                    f"({a['start']}–open) that is not the last entry"
                )
            if a["end"] >= b["start"]:
                raise ValueError(
                    f"account_periods.json: overlapping periods for {ticker!r}: "
                    f"{a['start']}–{a['end']} overlaps {b['start']}–{b['end']}"
                )
        parsed[ticker.strip().upper()] = parsed_periods

    return parsed


def save_account_periods(path, periods_dict):
    output = {}
    for ticker, periods in sorted(periods_dict.items()):
        output[ticker] = []
        for p in periods:
            entry = {"account": p["account"], "start": p["start"].strftime("%Y-%m-%d")}
            if p["end"] is not None:
                entry["end"] = p["end"].strftime("%Y-%m-%d")
            output[ticker].append(entry)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2)
        f.write("\n")


# ---------------------------------------------------------------------------
# security_map.json
# ---------------------------------------------------------------------------

def load_security_map(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Malformed security_map.json: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("security_map.json must be a JSON object.")
    return {k.strip().upper(): v for k, v in data.items()}


def save_security_map(path, mapping):
    with open(path, "w", encoding="utf-8") as f:
        json.dump({k: v for k, v in sorted(mapping.items())}, f, indent=2)
        f.write("\n")


# ---------------------------------------------------------------------------
# Account / security resolution
# ---------------------------------------------------------------------------

def resolve_account(txn, account_periods, default_account=None, verbose=False):
    ticker     = txn["ticker"].strip().upper()
    trade_date = txn["trade_date"]

    if account_periods:
        periods = account_periods.get(ticker)
        if periods:
            for p in periods:
                if p["start"] <= trade_date:
                    if p["end"] is None or trade_date <= p["end"]:
                        if verbose:
                            print(f"[ACCOUNT] {ticker} {trade_date}: "
                                  f"{p['start']}–{p['end'] or 'open'} -> {p['account']!r}")
                        return p["account"]
            if verbose:
                print(f"[ACCOUNT] {ticker} {trade_date}: no matching period")
            return None

    if default_account:
        if verbose:
            print(f"[ACCOUNT] {ticker} {trade_date}: using default {default_account!r}")
        return default_account

    return None


def resolve_security(txn, security_map, verbose=False):
    ticker = txn["ticker"].strip().upper()
    if security_map:
        name = security_map.get(ticker)
        if name:
            if verbose:
                print(f"[SECURITY] {ticker} -> {name!r}")
            return name
    if verbose:
        print(f"[SECURITY] {ticker}: no mapping, using ticker")
    return ticker


# ---------------------------------------------------------------------------
# Interactive setup
# ---------------------------------------------------------------------------

def interactive_account_setup(tickers):
    print("\nEnter the Quicken account name for each fund.")
    print("If a fund moved between accounts, enter one account for now.")
    print("You can add more periods to account_periods.json afterward.\n")
    periods = {}
    for ticker in sorted(tickers):
        while True:
            account = input(f"  Quicken account for {ticker}: ").strip()
            if account:
                break
            print("  (cannot be empty)")
        periods[ticker] = [{"account": account, "start": date(1900, 1, 1), "end": None}]
    return periods


def interactive_security_setup(tickers):
    print("\nEnter the Quicken security name for each ticker.")
    print("(Press Enter to use the ticker symbol as-is.)\n")
    mapping = {}
    for ticker in sorted(tickers):
        name = input(f"  Security name for {ticker} [{ticker}]: ").strip()
        mapping[ticker] = name if name else ticker
    return mapping
